"""API end-to-end: login -> upload -> extract -> export."""

import io
import os

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app

client = TestClient(app)
SAMPLES = os.path.join(os.path.dirname(__file__), "..", "samples")


def _token() -> str:
    r = client.post("/api/auth/login",
                    json={"username": "demo", "password": "demo1234"})
    assert r.status_code == 200
    return r.json()["access_token"]


def _headers() -> dict:
    return {"Authorization": f"Bearer {_token()}"}


def test_login_rejects_bad_password():
    r = client.post("/api/auth/login",
                    json={"username": "demo", "password": "wrong"})
    assert r.status_code == 401


def test_protected_routes_need_token():
    assert client.get("/api/documents").status_code == 401


def test_full_flow():
    h = _headers()
    with open(os.path.join(SAMPLES, "sample-1-rechnung-korrekt.pdf"), "rb") as f:
        pdf = f.read()
    r = client.post("/api/upload", files={"file": ("inv.pdf", pdf, "application/pdf")},
                    headers=h)
    assert r.status_code == 200
    doc_id = r.json()["id"]
    assert r.json()["extracted"] is False

    r = client.get("/api/documents", headers=h)
    assert any(d["id"] == doc_id for d in r.json()["documents"])

    r = client.post(f"/api/documents/{doc_id}/extract", json={}, headers=h)
    assert r.status_code == 200
    body = r.json()
    assert body["extraction"]["fields"]["invoice_number"]["value"] == "RE-2026-001"
    assert body["extraction"]["fields"]["total"]["confidence"] >= 0.9

    r = client.get(f"/api/documents/{doc_id}/export.xlsx", headers=h)
    assert r.status_code == 200
    assert "spreadsheetml" in r.headers["content-type"]
    wb = load_workbook(io.BytesIO(r.content))
    assert "Summary" in wb.sheetnames


def test_export_before_extract_is_409():
    h = _headers()
    with open(os.path.join(SAMPLES, "sample-2-invoice-english.pdf"), "rb") as f:
        pdf = f.read()
    r = client.post("/api/upload", files={"file": ("inv2.pdf", pdf, "application/pdf")},
                    headers=h)
    doc_id = r.json()["id"]
    r = client.get(f"/api/documents/{doc_id}/export.xlsx", headers=h)
    assert r.status_code == 409


def test_unknown_extractor_rejected():
    h = _headers()
    with open(os.path.join(SAMPLES, "sample-1-rechnung-korrekt.pdf"), "rb") as f:
        pdf = f.read()
    r = client.post("/api/upload", files={"file": ("inv.pdf", pdf, "application/pdf")},
                    headers=h)
    doc_id = r.json()["id"]
    r = client.post(f"/api/documents/{doc_id}/extract",
                    json={"extractor": "nope"}, headers=h)
    assert r.status_code == 400
