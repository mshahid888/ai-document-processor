"""Export produces a real, readable .xlsx workbook."""

import io

from openpyxl import load_workbook

from app.export import build_workbook
from app.extraction.base import ExtractionResult, FieldResult, LineItem
from app.validation import ValidationIssue


def _doc():
    ext = ExtractionResult(extractor_name="test")
    ext.fields["invoice_number"] = FieldResult("RE-1", 0.95, "regex:invoice_number")
    ext.fields["total"] = FieldResult(119.0, 0.9, "regex:total")
    ext.fields["missing_field"] = FieldResult(
        None, 0.0, "none", warning="not found (not guessed)")
    ext.line_items = [LineItem("Service", 2, 50.0, 100.0)]
    return {
        "filename": "sample.pdf",
        "uploaded_at": "2026-10-14T10:00:00",
        "extraction": ext,
        "issues": [ValidationIssue("warning", "totals_mismatch", "sums differ")],
    }


def test_workbook_structure():
    xlsx = build_workbook([_doc()])
    wb = load_workbook(io.BytesIO(xlsx))
    names = wb.sheetnames
    assert "Summary" in names
    assert any("fields" in n for n in names)
    assert any("items" in n for n in names)
    assert any("validation" in n for n in names)


def test_summary_row_values():
    xlsx = build_workbook([_doc()])
    wb = load_workbook(io.BytesIO(xlsx))
    ws = wb["Summary"]
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][0] == "File"  # header
    data = rows[1]
    assert data[0] == "sample.pdf"
    assert data[2] == "RE-1"
    assert data[5] == 119.0
    assert data[7] == 1  # one validation issue


def test_field_sheet_has_confidence_and_warning():
    xlsx = build_workbook([_doc()])
    wb = load_workbook(io.BytesIO(xlsx))
    ws = next(ws for ws in wb.worksheets if ws.title.endswith("fields"))
    rows = {r[0]: r for r in ws.iter_rows(values_only=True)}
    assert rows["invoice_number"][2] == 0.95
    assert rows["missing_field"][1] in (None, "")
    assert "not found" in str(rows["missing_field"][4])
